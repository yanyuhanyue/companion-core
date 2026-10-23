"""Native Linux Docker acceptance for the standalone file parser; no live user data.

Build deploy/file-parser.Dockerfile, inspect the local image's sha256 ID, then run:
    python scripts/smoke_file_parser.py --image sha256:<image-id> --verify-launcher-death
The host needs only Python's standard library and Docker. Fixtures, including a
real legacy DOC conversion, are generated inside the same pinned offline image.
"""

import argparse
import base64
import json
import os
import subprocess
import sys
import tempfile
import time
import uuid
from dataclasses import replace
from pathlib import Path

REPOSITORY = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPOSITORY))

from companion import file_parser  # noqa: E402

GENERATOR = r'''
import base64
import io
import json
import logging
import os
import resource
import subprocess
from pathlib import Path

import docx
import pypdf
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

logging.disable(logging.CRITICAL)
status = dict(line.split(":", 1) for line in Path("/proc/self/status").read_text().splitlines()
              if ":" in line)
root_mount = next(line.split() for line in Path("/proc/mounts").read_text().splitlines()
                  if line.split()[1] == "/")
assert os.getuid() == 10003
assert int(status["CapEff"].strip(), 16) == 0
assert status["NoNewPrivs"].strip() == "1"
assert "ro" in root_mount[3].split(",")
assert sorted(path.name for path in Path("/sys/class/net").iterdir()) == ["lo"]
assert resource.getrlimit(resource.RLIMIT_CPU) == (15, 15)
assert resource.getrlimit(resource.RLIMIT_FSIZE) == (134217728, 134217728)

def cgroup_value(unified, legacy):
    path = Path("/sys/fs/cgroup") / unified
    if not path.is_file():
        path = Path("/sys/fs/cgroup") / legacy
    return path.read_text().strip()

assert cgroup_value("memory.max", "memory/memory.limit_in_bytes") == "805306368"
assert cgroup_value("pids.max", "pids/pids.max") == "96"
cpu_max = Path("/sys/fs/cgroup/cpu.max")
if cpu_max.is_file():
    quota, period = cpu_max.read_text().split()
else:
    quota = cgroup_value("cpu/cpu.cfs_quota_us", "cpu,cpuacct/cpu.cfs_quota_us")
    period = cgroup_value("cpu/cpu.cfs_period_us", "cpu,cpuacct/cpu.cfs_period_us")
assert 0 < int(quota) <= int(period)

document = docx.Document()
document.add_paragraph("before table")
table = document.add_table(rows=1, cols=2)
table.cell(0, 0).text = "left cell"
table.cell(0, 1).text = "right cell"
document.add_paragraph("after table")
docx_stream = io.BytesIO()
document.save(docx_stream)

def pdf(texts, encrypted=False):
    writer = pypdf.PdfWriter()
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    for text in texts:
        page = writer.add_blank_page(width=300, height=400)
        if text is not None:
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
            stream = DecodedStreamObject()
            stream.set_data(b"BT /F1 12 Tf 10 300 Td (" + text.encode("ascii") + b") Tj ET")
            page[NameObject("/Contents")] = stream
    if encrypted:
        writer.encrypt("synthetic-fixture-password")
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()

scratch = Path("/scratch")
source = scratch / "fixture.docx"
source.write_bytes(docx_stream.getvalue())
profile = scratch / "generator-profile"
destination = scratch / "legacy"
destination.mkdir()
conversion = subprocess.run([
    "soffice", "-env:UserInstallation=" + profile.as_uri(), "--headless", "--nologo",
    "--nodefault", "--norestore", "--convert-to", "doc:MS Word 97", "--outdir",
    str(destination), str(source),
], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=25, check=False)
assert conversion.returncode == 0 and (destination / "fixture.doc").is_file()
legacy = (destination / "fixture.doc").read_bytes()
assert legacy.startswith(bytes.fromhex("d0cf11e0a1b11ae1"))
fixtures = {
    "text.txt": "中文合成文档\n\nsecond paragraph".encode(),
    "notes.md": b"# Synthetic heading\n\nmarkdown paragraph",
    "table.docx": docx_stream.getvalue(),
    "pages.pdf": pdf(["first page", "second page"]),
    "without-text.pdf": pdf([None]),
    "encrypted.pdf": pdf(["encrypted text"], encrypted=True),
    "legacy.doc": legacy,
}
assert sum(len(value) for value in fixtures.values()) < 4 * 1024 * 1024
print(json.dumps({
    "fixtures": {name: base64.b64encode(data).decode("ascii") for name, data in fixtures.items()},
    "versions": {"python_docx": docx.__version__, "pypdf": pypdf.__version__},
    "os_controls": {"uid": os.getuid(), "effective_capabilities": 0,
                    "no_new_privileges": True, "root_read_only": True,
                    "network_interfaces": ["lo"], "cpu_seconds": 15,
                    "memory_bytes": 805306368, "max_processes": 96},
}, separators=(",", ":")))
'''


def generate_fixtures(image, workspace, docker):
    limits = replace(file_parser.DEFAULT_LIMITS, wall_seconds=45)
    name = "companion-file-fixtures-" + uuid.uuid4().hex
    command = file_parser._container_argv(docker, image, name, "unused.txt", "", limits)
    command = command[:command.index(image) + 1] + [
        "--signal=KILL", str(limits.wall_seconds) + "s", "python", "-c", GENERATOR,
    ]
    scratch = workspace / ".local" / "test-temp"
    if not scratch.resolve().is_relative_to(workspace):
        raise ValueError("native fixture scratch must remain inside workspace")
    scratch.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="fixtures-", dir=scratch) as directory:
        try:
            completed = subprocess.run(
                command, shell=False, cwd=directory, capture_output=True, timeout=45, check=False,
                creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
            )
            if completed.returncode:
                detail = completed.stderr.decode("utf-8", errors="replace")[-2000:]
                raise RuntimeError("native fixture container failed: " + detail)
            if len(completed.stdout) > 8 * 1024 * 1024:
                raise RuntimeError("native fixture output exceeded limit")
            result = json.loads(completed.stdout)
            result["fixtures"] = {
                name: base64.b64decode(value, validate=True)
                for name, value in result["fixtures"].items()
            }
            return result
        except Exception:
            try:
                subprocess.run([docker, "rm", "--force", name], shell=False, cwd=directory,
                               stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                               timeout=5, check=False)
            except (OSError, subprocess.TimeoutExpired):
                pass
            raise


def verify_launcher_death(image, workspace, docker):
    """Kill only this smoke's launcher; the container must enforce its own deadline."""
    if sys.platform != "linux":
        raise RuntimeError("launcher SIGKILL acceptance requires Linux")
    limits = replace(file_parser.DEFAULT_LIMITS, wall_seconds=5)
    name, owner = "companion-file-" + uuid.uuid4().hex, uuid.uuid4().hex
    command = file_parser._container_argv(docker, image, name, "unused.txt", "", limits,
                                          owner_label=owner)
    # A trusted synthetic workload outlives the deadline and spawns a descendant.
    workload = ("import subprocess,time; "
                "subprocess.Popen(['python','-c','import time; time.sleep(600)']); "
                "time.sleep(600)")
    command = command[:command.index(image) + 1] + [
        "--signal=KILL", "5s", "python", "-c", workload,
    ]
    child = ("import json,subprocess,sys; "
             "p=subprocess.Popen(json.loads(sys.argv[1]),stdin=subprocess.DEVNULL,"
             "stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL); p.wait()")
    process = subprocess.Popen([sys.executable, "-I", "-c", child, json.dumps(command)],
                               shell=False, cwd=workspace, stdin=subprocess.DEVNULL,
                               stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

    def state():
        result = subprocess.run(
            [docker, "ps", "-a", "--filter", "label=com.companion.parser-host=" + owner,
             "--format", "{{.Names}} {{.State}}"], shell=False, cwd=workspace,
            capture_output=True, timeout=2, check=False,
        )
        if result.returncode:
            raise RuntimeError("cannot confirm native container state")
        rows = result.stdout.decode("ascii").splitlines()
        if any(not row.startswith(name + " ") for row in rows):
            raise RuntimeError("unexpected native acceptance container identity")
        return rows

    try:
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            if name + " running" in state():
                break
            if process.poll() is not None:
                raise RuntimeError("native launcher exited before running container was observed")
            time.sleep(0.05)
        else:
            raise RuntimeError("native acceptance container did not start")
        process.kill()  # Linux SIGKILL; no launcher finally block can clean up this container.
        process.wait(timeout=2)
        killed_at = time.monotonic()
        while time.monotonic() - killed_at < limits.wall_seconds + 5:
            if not state():
                return {"launcher_signal": "SIGKILL", "running_container_observed": True,
                        "container_removed": True, "inner_wall_seconds": limits.wall_seconds,
                        "seconds_after_launcher_death": round(time.monotonic() - killed_at, 3)}
            time.sleep(0.1)
        raise RuntimeError("container survived its independent wall deadline")
    finally:
        if process.poll() is None:
            process.kill()
            process.wait(timeout=2)
        # Only the random name created by this invocation is eligible for cleanup.
        subprocess.run([docker, "rm", "--force", name], shell=False, cwd=workspace,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=5,
                       check=False)


def main(argv=None):
    arguments = argparse.ArgumentParser(description=__doc__)
    arguments.add_argument("--image", required=True, help="local sha256 ID or repository@sha256")
    arguments.add_argument("--workspace", type=Path,
                           default=REPOSITORY / ".local/file-parser-native")
    arguments.add_argument("--docker", default="docker")
    arguments.add_argument("--verify-launcher-death", action="store_true")
    args = arguments.parse_args(argv)
    workspace = args.workspace.resolve()
    if not workspace.is_relative_to(REPOSITORY):
        arguments.error("workspace must remain inside the repository")
    workspace.mkdir(parents=True, exist_ok=True)
    generated = generate_fixtures(args.image, workspace, args.docker)
    cases = {}

    def parse(name, expected, *, limits=file_parser.DEFAULT_LIMITS, label=None):
        data = generated["fixtures"][name]
        result = file_parser.parse_isolated(data, name, image=args.image, workspace=workspace,
                                            limits=limits, docker=args.docker)
        case = label or name
        if result["status"] != expected:
            raise RuntimeError(f"{case}: {result['status']}/{result['error_code']}")
        if result["isolation"] != "docker_network_none":
            raise RuntimeError(f"{case}: actual isolated parser did not complete")
        cases[case] = {"status": result["status"], "issues": result["issues"],
                       "error_code": result["error_code"], "records": len(result["records"])}
        return result

    text = parse("text.txt", "SUCCESS")
    assert text["records"][0]["source"] == {"line_start": 1, "line_end": 1}
    assert text["text"].startswith("中文合成文档")
    markdown = parse("notes.md", "SUCCESS")
    assert markdown["text"].startswith("# Synthetic heading")
    document = parse("table.docx", "SUCCESS")
    assert [row["text"] for row in document["records"]] == [
        "before table", "left cell", "right cell", "after table",
    ]
    assert document["records"][1]["source"]["table_path"] == [1, 1, 1]
    pdf = parse("pages.pdf", "SUCCESS")
    assert [row["source"] for row in pdf["records"]] == [{"page": 1}, {"page": 2}]
    assert pdf["text"] == "first page\nsecond page"
    legacy = parse("legacy.doc", "PARTIAL")
    assert "DOC_CONVERTED_TO_DOCX" in legacy["issues"]
    assert "left cell" in legacy["text"] and "after table" in legacy["text"]
    without_text = parse("without-text.pdf", "PARTIAL")
    assert without_text["issues"] == ["PAGE_WITHOUT_TEXT_LAYER"]
    assert without_text["text"] == "" and without_text["ocr"] == "not_attempted"
    encrypted = parse("encrypted.pdf", "ERROR")
    assert encrypted["error_code"] == "ENCRYPTED_DOCUMENT"
    bounded_text = parse("text.txt", "PARTIAL", label="text-character-limit",
                         limits=replace(file_parser.DEFAULT_LIMITS, max_text_chars=3))
    assert bounded_text["issues"] == ["TEXT_LIMIT"] and len(bounded_text["text"]) == 3
    bounded_pages = parse("pages.pdf", "PARTIAL", label="pdf-page-limit",
                          limits=replace(file_parser.DEFAULT_LIMITS, max_pages=1))
    assert bounded_pages["issues"] == ["PAGE_LIMIT"] and len(bounded_pages["records"]) == 1
    bounded_zip = parse("table.docx", "ERROR", label="docx-expanded-limit",
                        limits=replace(file_parser.DEFAULT_LIMITS, max_expanded_bytes=64))
    assert bounded_zip["error_code"] == "ZIP_EXPANSION_LIMIT"
    launcher_death = verify_launcher_death(args.image, workspace, args.docker) if (
        args.verify_launcher_death
    ) else None
    print(json.dumps({"status": "PASS", "image": args.image, "cases": cases,
                      "versions": generated["versions"],
                      "os_controls": generated["os_controls"],
                      "launcher_death": launcher_death}, ensure_ascii=True))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(json.dumps({"status": "FAIL", "error": str(error) or type(error).__name__}))
        raise SystemExit(1) from None
